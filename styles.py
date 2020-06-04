from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from tools import create_paragraph_style
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml
from tools import get_cells_coordinate


def merge(table, value):
    # row-col/row-col!
    for elt in value.split('!'):
        begin_cell, end_cell = elt.split('/')
        begin_row, begin_column = begin_cell.split('-')
        end_row, end_column = end_cell.split('-')

        table.cell(int(begin_row), int(begin_column)).merge(table.cell(int(end_row), int(end_column)))


def global_paragraph_style(paragraph, value):
    paragraph.style = value


def caption_style(paragraph, value):
    if paragraph.text != '':
        paragraph.style = value


def paragraph_font_size(paragraph, value):
    new_style = create_paragraph_style(paragraph)
    new_style.font.size = Pt(int(value))
    paragraph.style = new_style


def paragraph_font_color(paragraph, value):
    value = value[1:]
    while len(value) < 6:
        value = '0'+value
    new_style = create_paragraph_style(paragraph)
    new_style.font.color.rgb = RGBColor(int(value[:2], 16), int(value[2:4], 16), int(value[4:], 16))
    paragraph.style = new_style


def cells_background_color(table, value):
    # value = !11-15/*-12-15#00000F
    for elt in value.split('!'):
        area, color = elt.split('#')
        rows, columns = get_cells_coordinate(table, area)
        for i in rows:
            for j in columns:
                shading_elm_1 = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color))
                table.cell(i,j)._tc.get_or_add_tcPr().append(shading_elm_1)


def cells_alignement(table, value):
    # value = !11-15/*-12-15#CENTER
    for elt in value.split('!'):
        area, alignment = elt.split('#')
        rows, columns = get_cells_coordinate(table, area)
        for i in rows:
            for j in columns:
                for paragraph in table.cell(i, j).paragraphs:
                    paragraph_align(paragraph, alignment)


def global_table_style(table, value):
    table.style = value


def paragraph_indentation(paragraph, value):
    paragraph.paragraph_format.left_indent = Pt(int(value))


def table_indentation(table, value):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph_indentation(paragraph, value)


def paragraph_align(paragraph, value):
    if type(paragraph) != list:
        paragraphs = [paragraph]
    else:
        paragraphs = paragraph
    for elt in paragraphs:
        elt.alignment = WD_ALIGN_PARAGRAPH.__dict__[value]


def set_autofit(table, value):
    table.autofit = True
    table.allow_autofit = True
    table._tblPr.xpath("./w:tblW")[0].attrib["{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type"] = "auto"
    for i in range(len(table.rows)):
        for j in range(len((table.rows[0].cells))):
            table.cell(i, j)._tc.tcPr.tcW.type = 'auto'
            table.cell(i, j)._tc.tcPr.tcW.w = 0
    return table
